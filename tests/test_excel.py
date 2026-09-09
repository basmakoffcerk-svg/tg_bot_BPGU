import datetime
import io

import openpyxl
import pytest
from sqlalchemy.ext.asyncio import AsyncSession

from services.excel_generator import generate_attendance_excel


@pytest.mark.asyncio
async def test_excel_generation(db_session: AsyncSession, seed_test_data):
    today = datetime.date.today()
    monday = today - datetime.timedelta(days=today.weekday())

    excel_buf, filename, total_pairs, total_absent_hours = await generate_attendance_excel(
        session=db_session, date_from=monday, date_to=today
    )

    assert isinstance(excel_buf, io.BytesIO)
    assert filename.endswith(".xlsx")
    assert "240326" in filename

    # Load and inspect generated workbook
    wb = openpyxl.load_workbook(excel_buf)
    ws = wb.active
    assert ws.title == "Рапортичка 240326"

    # Header assertions
    assert "БАШКИРСКИЙ ГОСУДАРСТВЕННЫЙ ПЕДАГОГИЧЕСКИЙ УНИВЕРСИТЕТ" in ws["A1"].value
    assert "240326" in ws["A3"].value

    # Check student rows
    cell_student = ws.cell(row=7, column=2).value
    assert cell_student is not None
    assert len(cell_student) > 0
