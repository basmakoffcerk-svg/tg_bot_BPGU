import datetime
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from models import (
    Attendance,
    AttendanceStatusEnum,
    PairsRegistry,
    ScheduleSlot,
    Student,
    Subject,
)

# Pre-compiled styles for performance
FONT_TITLE = Font(name="Calibri", size=14, bold=True)
FONT_SUBTITLE = Font(name="Calibri", size=11, bold=True)
FONT_HEADER = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
FONT_REGULAR = Font(name="Calibri", size=10)
FONT_BOLD = Font(name="Calibri", size=10, bold=True)

FILL_HEADER = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
FILL_SUBHEADER = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
FILL_TOTAL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

FILL_ABSENT_UNEXCUSED = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")  # Light Red
FILL_ABSENT_EXCUSED = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")  # Light Green
FILL_LATE = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # Light Yellow

ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

THIN_BORDER_SIDE = Side(border_style="thin", color="000000")
DOUBLE_BORDER_BOTTOM = Side(border_style="double", color="000000")

BORDER_CELL = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=THIN_BORDER_SIDE)
BORDER_TOTAL = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=DOUBLE_BORDER_BOTTOM)


def sanitize_excel_cell(val: str) -> str:
    """Prevents CSV/Excel formula injection (CWE-1236)."""
    if isinstance(val, str) and len(val) > 0 and val[0] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + val
    return val


async def generate_attendance_excel(
    session: AsyncSession, date_from: datetime.date, date_to: datetime.date
) -> tuple[io.BytesIO, str, int, int]:
    """
    Generates official university attendance report (.xlsx) for group 240326.
    Returns: (BytesIO buffer, filename, total_pairs, total_absent_hours)
    """
    if date_from > date_to:
        raise ValueError("date_from must be <= date_to")

    # 1. Fetch students sorted by full_name
    res_students = await session.execute(select(Student).order_by(Student.full_name))
    students = res_students.scalars().all()

    # 2. Fetch pairs registry in date range with slot and subject
    res_pairs = await session.execute(
        select(PairsRegistry, ScheduleSlot, Subject)
        .join(ScheduleSlot, PairsRegistry.slot_id == ScheduleSlot.id)
        .join(Subject, ScheduleSlot.subject_id == Subject.id)
        .where(PairsRegistry.calendar_date >= date_from, PairsRegistry.calendar_date <= date_to)
        .order_by(PairsRegistry.calendar_date, ScheduleSlot.pair_number)
    )
    pair_rows = res_pairs.all()

    # 3. Fetch attendance records in date range
    pair_ids = [p[0].id for p in pair_rows]
    attendances_map: dict[tuple[int, int], Attendance] = {}
    if pair_ids:
        res_att = await session.execute(select(Attendance).where(Attendance.pair_id.in_(pair_ids)))
        for att in res_att.scalars().all():
            attendances_map[(att.pair_id, att.student_id)] = att

    # 4. Create Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Рапортичка 240326"
    ws.views.sheetView[0].showGridLines = True

    # Calculate total columns: 3 fixed + len(pair_rows) + 3 summary columns
    col_total = 3 + len(pair_rows) + 3
    max_col_letter = get_column_letter(max(col_total, 6))

    # Header Title Block (dynamically sized)
    ws.merge_cells(f"A1:{max_col_letter}1")
    ws["A1"] = "БАШКИРСКИЙ ГОСУДАРСТВЕННЫЙ ПЕДАГОГИЧЕСКИЙ УНИВЕРСИТЕТ ИМ. М. АКМУЛЛЫ"
    ws["A1"].font = FONT_TITLE
    ws["A1"].alignment = ALIGN_CENTER

    ws.merge_cells(f"A2:{max_col_letter}2")
    ws["A2"] = "Институт физики, математики, цифровых и нанотехнологий"
    ws["A2"].font = FONT_SUBTITLE
    ws["A2"].alignment = ALIGN_CENTER

    ws.merge_cells(f"A3:{max_col_letter}3")
    ws["A3"] = "ЖУРНАЛ УЧЕТА ПОСЕЩАЕМОСТИ (РАПОРТИЧКА) | ГРУППА: 240326 «МАТИНФ»"
    ws["A3"].font = FONT_SUBTITLE
    ws["A3"].alignment = ALIGN_CENTER

    ws.merge_cells(f"A4:{max_col_letter}4")
    ws["A4"] = f"Отчетный период: {date_from.strftime('%d.%m.%Y')} — {date_to.strftime('%d.%m.%Y')}"
    ws["A4"].font = FONT_REGULAR
    ws["A4"].alignment = ALIGN_CENTER

    # Table Header Row
    row_idx = 6
    ws.cell(row=row_idx, column=1, value="№").font = FONT_HEADER
    ws.cell(row=row_idx, column=1).fill = FILL_HEADER
    ws.cell(row=row_idx, column=1).alignment = ALIGN_CENTER
    ws.cell(row=row_idx, column=1).border = BORDER_CELL

    ws.cell(row=row_idx, column=2, value="ФИО Студента").font = FONT_HEADER
    ws.cell(row=row_idx, column=2).fill = FILL_HEADER
    ws.cell(row=row_idx, column=2).alignment = ALIGN_CENTER
    ws.cell(row=row_idx, column=2).border = BORDER_CELL

    ws.cell(row=row_idx, column=3, value="Подгр.").font = FONT_HEADER
    ws.cell(row=row_idx, column=3).fill = FILL_HEADER
    ws.cell(row=row_idx, column=3).alignment = ALIGN_CENTER
    ws.cell(row=row_idx, column=3).border = BORDER_CELL

    col_idx = 4
    pair_col_mapping = []
    for pair_obj, slot_obj, subj_obj in pair_rows:
        col_title = f"{pair_obj.calendar_date.strftime('%d.%m')}\n№{slot_obj.pair_number}\n{subj_obj.title[:10]}"
        cell = ws.cell(row=row_idx, column=col_idx, value=col_title)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_CELL
        pair_col_mapping.append((col_idx, pair_obj, slot_obj))
        col_idx += 1

    # Summary columns
    col_unexcused = col_idx
    c_un = ws.cell(row=row_idx, column=col_unexcused, value="Н (часов)")
    c_un.font = FONT_HEADER
    c_un.fill = FILL_HEADER
    c_un.alignment = ALIGN_CENTER
    c_un.border = BORDER_CELL

    col_excused = col_idx + 1
    c_ex = ws.cell(row=row_idx, column=col_excused, value="У (часов)")
    c_ex.font = FONT_HEADER
    c_ex.fill = FILL_HEADER
    c_ex.alignment = ALIGN_CENTER
    c_ex.border = BORDER_CELL

    col_total = col_idx + 2
    c_tot = ws.cell(row=row_idx, column=col_total, value="Всего пропусков (ч)")
    c_tot.font = FONT_HEADER
    c_tot.fill = FILL_HEADER
    c_tot.alignment = ALIGN_CENTER
    c_tot.border = BORDER_CELL

    ws.row_dimensions[row_idx].height = 40

    # Data Rows
    row_idx += 1
    total_group_unexcused_hours = 0
    total_group_excused_hours = 0

    for i, student in enumerate(students, 1):
        ws.cell(row=row_idx, column=1, value=i).alignment = ALIGN_CENTER
        ws.cell(row=row_idx, column=1).border = BORDER_CELL
        ws.cell(row=row_idx, column=1).font = FONT_REGULAR

        ws.cell(row=row_idx, column=2, value=sanitize_excel_cell(student.full_name)).alignment = ALIGN_LEFT
        ws.cell(row=row_idx, column=2).border = BORDER_CELL
        ws.cell(row=row_idx, column=2).font = FONT_REGULAR

        ws.cell(row=row_idx, column=3, value=student.subgroup).alignment = ALIGN_CENTER
        ws.cell(row=row_idx, column=3).border = BORDER_CELL
        ws.cell(row=row_idx, column=3).font = FONT_REGULAR

        student_unexcused_pairs = 0
        student_excused_pairs = 0

        for col, pair_obj, slot_obj in pair_col_mapping:
            cell = ws.cell(row=row_idx, column=col)
            cell.border = BORDER_CELL
            cell.alignment = ALIGN_CENTER
            cell.font = FONT_REGULAR

            # Check if this slot belongs to student subgroup
            if slot_obj.subgroup != 0 and slot_obj.subgroup != student.subgroup:
                cell.value = "—"
                cell.font = Font(name="Calibri", size=9, color="808080")
                continue

            att = attendances_map.get((pair_obj.id, student.id))
            if not att:
                # Absent by default if pair is logged
                cell.value = "Н"
                cell.fill = FILL_ABSENT_UNEXCUSED
                student_unexcused_pairs += 1
            elif att.status in (AttendanceStatusEnum.PRESENT.value, AttendanceStatusEnum.MANUAL_CONFIRM.value):
                cell.value = "·"
            elif att.status == AttendanceStatusEnum.LATE.value:
                cell.value = "О"
                cell.fill = FILL_LATE
            elif att.status == AttendanceStatusEnum.ABSENT_EXCUSED.value:
                cell.value = "У"
                cell.fill = FILL_ABSENT_EXCUSED
                student_excused_pairs += 1
            else:  # ABSENT_UNEXCUSED
                cell.value = "Н"
                cell.fill = FILL_ABSENT_UNEXCUSED
                student_unexcused_pairs += 1

        # Calculate academic hours (1 pair = 2 hours)
        unexcused_hours = student_unexcused_pairs * 2
        excused_hours = student_excused_pairs * 2
        total_student_hours = unexcused_hours + excused_hours

        total_group_unexcused_hours += unexcused_hours
        total_group_excused_hours += excused_hours

        c_u = ws.cell(row=row_idx, column=col_unexcused, value=unexcused_hours)
        c_u.alignment = ALIGN_CENTER
        c_u.border = BORDER_CELL
        c_u.font = FONT_BOLD

        c_e = ws.cell(row=row_idx, column=col_excused, value=excused_hours)
        c_e.alignment = ALIGN_CENTER
        c_e.border = BORDER_CELL
        c_e.font = FONT_BOLD

        c_t = ws.cell(row=row_idx, column=col_total, value=total_student_hours)
        c_t.alignment = ALIGN_CENTER
        c_t.border = BORDER_CELL
        c_t.font = FONT_BOLD

        row_idx += 1

    # Total Summary Row
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=3)
    tot_label = ws.cell(row=row_idx, column=1, value="ИТОГО ПО ГРУППЕ (часов)")
    tot_label.font = FONT_BOLD
    tot_label.alignment = ALIGN_RIGHT
    tot_label.fill = FILL_TOTAL

    for c in range(1, col_total + 1):
        ws.cell(row=row_idx, column=c).border = BORDER_TOTAL
        ws.cell(row=row_idx, column=c).fill = FILL_TOTAL

    tot_un = ws.cell(row=row_idx, column=col_unexcused, value=total_group_unexcused_hours)
    tot_un.font = FONT_BOLD
    tot_un.alignment = ALIGN_CENTER

    tot_ex = ws.cell(row=row_idx, column=col_excused, value=total_group_excused_hours)
    tot_ex.font = FONT_BOLD
    tot_ex.alignment = ALIGN_CENTER

    tot_all = ws.cell(row=row_idx, column=col_total, value=total_group_unexcused_hours + total_group_excused_hours)
    tot_all.font = FONT_BOLD
    tot_all.alignment = ALIGN_CENTER

    # Auto-adjust column widths
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 8
    for c in range(4, col_idx):
        col_letter = get_column_letter(c)
        ws.column_dimensions[col_letter].width = 12
    ws.column_dimensions[get_column_letter(col_unexcused)].width = 12
    ws.column_dimensions[get_column_letter(col_excused)].width = 12
    ws.column_dimensions[get_column_letter(col_total)].width = 18

    # Save to memory buffer
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Рапортичка_240326_{date_from.strftime('%d.%m')}_{date_to.strftime('%d.%m.%Y')}.xlsx"
    total_absent_hours = total_group_unexcused_hours + total_group_excused_hours

    return output, filename, len(pair_rows), total_absent_hours
